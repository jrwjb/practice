import datetime
from docx import Document

def record_doc():
    today = datetime.date.today()
    rec_doc = Document('C:\\Users\\jbwang\\Desktop\\Labelfree项目分析流转单及原始记录表.docx')
    tables = rec_doc.tables
    project_name = tables[0].cell(0, 2).text
    client = tables[0].cell(1, 2).text
    project_id = tables[0].cell(2, 2).text
    species = tables[0].cell(3, 2).text
    database = tables[0].cell(4, 2).text
    database_file = tables[0].cell(5, 2).text
    sample_info = tables[0].cell(9, 2)
    # groupvs_num = tables[0].cell(8, 2).text
    # repeat_num = tables[0].cell(9, 2).text
    # samples_num = tables[0].cell(10, 2).text
    # sample_info = tables[0].cell(11, 1)  # 样本表信息
    # group_vs = tables[0].cell(12, 3).text
    # go_kegg = tables[0].cell(13, 3).text
    # go_kegg_enrich = tables[0].cell(14, 3).text
    # cluster = tables[0].cell(15, 3).text
    # ppi = tables[0].cell(16, 3).text
    # motif = tables[0].cell(17, 3).text
    return (project_name, client, project_id, today, species, database, database_file, sample_info)